from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

wb = Workbook()
ws = wb.active
ws.title = "Vendor Master"

headers = [
    "S. No.",
    "Vendor company Name",
    "For Sector",
    "Office/Registered address",
    "Sign authority Name",
    "contact No.",
    "Email id",
    "Agreement Type RCM- NO GST ON BILLS / FCM- GST ON BILLS",
    "Available Vehicle Types",
    "Count of Own Vehicle",
]

widths = [8, 28, 16, 36, 22, 16, 26, 42, 24, 18]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

thin = Border(
    left=Side(style="thin", color="334155"),
    right=Side(style="thin", color="334155"),
    top=Side(style="thin", color="334155"),
    bottom=Side(style="thin", color="334155"),
)
header_fill = PatternFill("solid", fgColor="0F766E")
header_font = Font(bold=True, color="FFFFFF", size=11, name="Calibri")
cell_font = Font(size=11, name="Calibri")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left = Alignment(horizontal="left", vertical="center", wrap_text=True)

ws.row_dimensions[1].height = 48

for col, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col, value=h)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = center
    cell.border = thin

for r in range(2, 27):
    ws.row_dimensions[r].height = 22
    for c in range(1, 11):
        cell = ws.cell(row=r, column=c, value=(r - 1) if c == 1 else None)
        cell.font = cell_font
        cell.border = thin
        cell.alignment = center if c in (1, 6, 10) else left

dv = DataValidation(
    type="list",
    formula1='"RCM - NO GST ON BILLS,FCM - GST ON BILLS"',
    allow_blank=True,
)
dv.error = "Select RCM or FCM"
dv.errorTitle = "Agreement Type"
dv.prompt = "Select agreement type"
dv.promptTitle = "Agreement Type"
ws.add_data_validation(dv)
dv.add("H2:H26")

ws.merge_cells("A28:J28")
ws.merge_cells("A29:J29")
ws.merge_cells("A30:J30")

note1 = ws.cell(
    row=28,
    column=1,
    value="Business Operated as Proprietorship / Partnership / Company : _______________________________",
)
note1.font = Font(bold=True, size=11, name="Calibri")
note1.alignment = Alignment(vertical="center")

note2 = ws.cell(
    row=29,
    column=1,
    value="(Partnership deed to be attached in case Partnership)",
)
note2.font = Font(italic=True, size=10, color="475569", name="Calibri")

note3 = ws.cell(
    row=30,
    column=1,
    value="Agreement Type: RCM = NO GST ON BILLS  |  FCM = GST ON BILLS",
)
note3.font = Font(size=10, color="0F766E", name="Calibri")

ws.row_dimensions[28].height = 24
ws.row_dimensions[29].height = 20
ws.row_dimensions[30].height = 20

ws.freeze_panes = "A2"
ws.auto_filter.ref = "A1:J26"

out = r"e:\project\vmcindia\Vendor_Company_Master.xlsx"
wb.save(out)
print(out)
