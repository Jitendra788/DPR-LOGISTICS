"""Fill BA Selection Criteria + Vendor Company Master from DPR Logistics project data."""
from datetime import date
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, Border, Side

BA_PATH = r"c:\Users\jiten\Downloads\BA Selection Criteria - V2.1.xlsx"
BA_OUT = r"e:\project\vmcindia\BA_Selection_Criteria_DPR_Logistics.xlsx"
VENDOR_PATH = r"e:\project\vmcindia\Vendor_Company_Master.xlsx"
VENDOR_OUT = r"e:\project\vmcindia\Vendor_Company_Master_Filled.xlsx"
VENDOR_DOWNLOADS = r"c:\Users\jiten\Downloads\Vendor_Company_Master_Filled.xlsx"
BA_DOWNLOADS = r"c:\Users\jiten\Downloads\BA_Selection_Criteria_DPR_Logistics.xlsx"

# --- Fill BA Selection Criteria (GUI sheet) ---
wb = load_workbook(BA_PATH)
ws = wb["GUI"]

ws["E8"] = "DPR Logistics"
ws["E9"] = "DPR-BA-001"
ws["E10"] = 9371662142
ws["E11"] = "BNLPK2073C"
ws["E12"] = "27BNLPK2073C1Z5"

# Questionnaire answers (exact option text from Options Weight tables)
answers = {
    15: "4. Sole Proprietorship",  # Type of Company (individual PAN)
    16: "2. 40 - 50 Yrs",  # MD/Partner age
    17: "3. Greater than 10 years",  # Logistics experience (founded 2014, ops since earlier)
    18: "1. Total Logistic Service",  # Part load, FTL, trailer, container, warehouse
    19: "Only MNC's/Pvt Co",  # Indo Count, Kirloskar, textile exporters etc.
    20: "< 50 Cr",  # Group turnover
    21: "< 30 Cr",  # Transport turnover
    22: "National",  # Pan-India lanes
    23: "< 20",  # Owned trucks
    24: "< 100",  # Attached / hired trucks
    25: "Management & inter office connectivity",  # Full ERP (LR, bill, GPS track)
    26: "Yes",  # GPS / live tracking
    27: "System generated MIS",  # MIS / GST / outstanding reports
    28: "Yes",  # Driver training
    29: "Yes",  # Driver insurance
    30: "Yes",  # Driver welfare
}
for row, value in answers.items():
    ws[f"E{row}"] = value

ws["E33"] = date.today().strftime("%d-%b-%Y")
ws["E34"] = "Procurement / Ops Review — DPR Logistics"

subjective = {
    37: "Books and GST billing maintained digitally via DPR Logistics ERP with LR, tax invoice and money receipt trail.",
    38: "Transparent freight billing with GST invoice (FCM) and clear party ledger / outstanding reports.",
    39: "Focused on long-term contract lanes across Maharashtra, Gujarat and pan-India industrial cargo.",
    40: "Already running digital LR booking, tracking, billing, POD and money receipt on dprlogistics.in ERP.",
    41: "Expanding digital customer portal, more lanes and attached fleet capacity.",
    42: "Strong base in Kolhapur MIDC with national coverage and enterprise customers.",
    43: "Services: part load, FTL, trailer, container, warehousing, GC/LR online tracking, GST billing.",
}
for row, value in subjective.items():
    ws[f"E{row}"] = value

wb.save(BA_OUT)
wb.save(BA_DOWNLOADS)
print("BA filled:", BA_OUT)
print("BA copy:", BA_DOWNLOADS)

# --- Fill Vendor Company Master ---
vendors = [
    {
        "name": "DPR Logistics",
        "sector": "Logistics / Transport",
        "address": "Shree Mahalaxmi Petrol Pump, 5 Star MIDC Road, Kagal, Dist. Kolhapur 416216",
        "signatory": "Proprietor / Authorized Signatory",
        "contact": "9356259949 / 9371662142",
        "email": "dprlogistics2142@gmail.com",
        "agreement": "FCM - GST ON BILLS",
        "vehicles": "Truck, Trailer, Container, Part Load / FTL",
        "own_count": "Own + Attached fleet",
        "business": "Proprietorship",
    },
    {
        "name": "RAJ TRANSPORT",
        "sector": "Transport Broker",
        "address": "Transport Nagar, Surat",
        "signatory": "Authorized Signatory",
        "contact": "9825001111",
        "email": "",
        "agreement": "RCM - NO GST ON BILLS",
        "vehicles": "Truck / Trailer (Attached)",
        "own_count": "",
        "business": "Proprietorship",
    },
    {
        "name": "SHREE LOGISTICS",
        "sector": "Transport Broker",
        "address": "Narol, Ahmedabad",
        "signatory": "Authorized Signatory",
        "contact": "9825002222",
        "email": "",
        "agreement": "RCM - NO GST ON BILLS",
        "vehicles": "Truck / Trailer (Attached)",
        "own_count": "",
        "business": "Proprietorship",
    },
    {
        "name": "IOCL Diesel Pump",
        "sector": "Fuel Vendor",
        "address": "NH-48, Surat",
        "signatory": "Branch Manager",
        "contact": "0261-2451001",
        "email": "",
        "agreement": "FCM - GST ON BILLS",
        "vehicles": "N/A",
        "own_count": "N/A",
        "business": "Company",
    },
    {
        "name": "National Insurance",
        "sector": "Insurance Vendor",
        "address": "CG Road, Ahmedabad",
        "signatory": "Branch Manager",
        "contact": "079-26581010",
        "email": "",
        "agreement": "FCM - GST ON BILLS",
        "vehicles": "N/A",
        "own_count": "N/A",
        "business": "Company",
    },
    {
        "name": "CITCO Roadways",
        "sector": "Transport Partner",
        "address": "Pan-India lane partner",
        "signatory": "Authorized Signatory",
        "contact": "",
        "email": "",
        "agreement": "FCM - GST ON BILLS",
        "vehicles": "Truck / Trailer",
        "own_count": "",
        "business": "Company",
    },
    {
        "name": "Welco Logistics",
        "sector": "Transport Partner",
        "address": "Pan-India lane partner",
        "signatory": "Authorized Signatory",
        "contact": "",
        "email": "",
        "agreement": "FCM - GST ON BILLS",
        "vehicles": "Truck / Trailer",
        "own_count": "",
        "business": "Company",
    },
    {
        "name": "MANMOHAN PRASAD PATHAK",
        "sector": "Transport Vendor",
        "address": "Transport Nagar",
        "signatory": "Manmohan Prasad Pathak",
        "contact": "9825004444",
        "email": "",
        "agreement": "RCM - NO GST ON BILLS",
        "vehicles": "Truck",
        "own_count": "",
        "business": "Proprietorship",
    },
    {
        "name": "SONU YADAV",
        "sector": "Transport Vendor",
        "address": "Jalgaon",
        "signatory": "Sonu Yadav",
        "contact": "9825005555",
        "email": "",
        "agreement": "RCM - NO GST ON BILLS",
        "vehicles": "Truck",
        "own_count": "",
        "business": "Proprietorship",
    },
    {
        "name": "PAWAN DOCTOR",
        "sector": "Transport Vendor",
        "address": "Surat",
        "signatory": "Pawan Doctor",
        "contact": "9825006666",
        "email": "",
        "agreement": "RCM - NO GST ON BILLS",
        "vehicles": "Truck",
        "own_count": "",
        "business": "Proprietorship",
    },
    {
        "name": "RAKESH KUMAR",
        "sector": "Transport Vendor",
        "address": "Ahmedabad",
        "signatory": "Rakesh Kumar",
        "contact": "9825007777",
        "email": "",
        "agreement": "RCM - NO GST ON BILLS",
        "vehicles": "Truck",
        "own_count": "",
        "business": "Proprietorship",
    },
]

vwb = load_workbook(VENDOR_PATH)
vws = vwb.active

thin = Border(
    left=Side(style="thin", color="334155"),
    right=Side(style="thin", color="334155"),
    top=Side(style="thin", color="334155"),
    bottom=Side(style="thin", color="334155"),
)
cell_font = Font(size=11, name="Calibri")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left = Alignment(horizontal="left", vertical="center", wrap_text=True)

# Clear previous sample rows 2-26 then fill
for r in range(2, 27):
    for c in range(1, 11):
        vws.cell(row=r, column=c).value = None

for i, v in enumerate(vendors, start=1):
    r = i + 1
    row_vals = [
        i,
        v["name"],
        v["sector"],
        v["address"],
        v["signatory"],
        v["contact"],
        v["email"],
        v["agreement"],
        v["vehicles"],
        v["own_count"],
    ]
    for c, val in enumerate(row_vals, 1):
        cell = vws.cell(row=r, column=c, value=val)
        cell.font = cell_font
        cell.border = thin
        cell.alignment = center if c in (1, 6, 10) else left

# Business type note — fill from primary company
vws["A28"] = (
    "Business Operated as Proprietorship / Partnership / Company : Proprietorship (DPR Logistics)"
)
vws["A29"] = "(Partnership deed to be attached in case Partnership)"
vws["A30"] = (
    "Agreement Type: RCM = NO GST ON BILLS  |  FCM = GST ON BILLS  |  Source: DPR Logistics ERP master / seed vendors"
)

vwb.save(VENDOR_OUT)
vwb.save(VENDOR_DOWNLOADS)
print("Vendor filled:", VENDOR_OUT)
print("Vendor copy:", VENDOR_DOWNLOADS)
print(f"Vendors rows: {len(vendors)}")
